from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(__file__).resolve().parents[1] / "artifacts" / "Prompts_GPT_Image_2_The_Unwritten_Road.docx"

SCENES = [
    (1, "La tinta todavía está fresca", "Presentar el llamado, la incertidumbre y el faro como destino desde el primer momento.",
     "Ultra-photorealistic cinematic epic fantasy, wide 16:9 composition. Early dawn inside a modest stone-and-timber traveler’s room. [PROTAGONIST] stands near an open wooden door, seen from behind and slightly in profile, holding an unsealed cream letter and a torn hand-drawn map. The map clearly suggests a road, an abandoned travelers’ station, a coastline and a remote lighthouse, but contains no readable text. Fresh black ink glistens on the paper. Beyond the doorway, across misty fields, an abandoned stone waystation and its bell tower are visible; the bell is moving although no person is there. A soft wind enters the room and lifts the corner of the letter. Emotional focus: the uneasy feeling of having been selected by someone unknown, mixed with curiosity strong enough to begin. Natural blue dawn light, subtle warm candle remnant, grounded medieval materials, realistic skin, cloth and paper. Keep the protagonist small enough that the mystery of the road dominates the frame. No other people, no visible magic beam, no text, no logos."),
    (2, "El comerciante que reconoció la carta", "Revelar la ceremonia del flautista sin explicar su motivo y mantener al protagonista en el camino principal.",
     "Ultra-photorealistic cinematic epic fantasy, wide 16:9 composition. A medieval night market beside the abandoned travelers’ station is closing under a rising full moon. Merchants lower canvas awnings, extinguish lanterns and pack wooden crates. In the foreground, [PROTAGONIST] holds the same unsealed letter while an older merchant with tired eyes suddenly leans forward in recognition, animated but not threatening. Behind him, a few details imply that a masked flutist performed here at sunset: an empty circular space, scattered white flower petals, a low stool and a faint memory-like silhouette reflected in a brass bowl, never a literal ghost. The merchant gestures toward the distant coastal road and a tiny lighthouse glow beyond the hills. Emotional focus: the tension between receiving useful information and realizing that the mystery is larger than expected. Warm amber lanterns against deep blue moonlight, realistic fabric, wood and faces. No transaction, no new object handed over, no companion joining the protagonist, no readable signs, no text or logos."),
    (3, "El puente roto", "Mostrar un obstáculo común donde diferentes estrategias caben en una sola composición.",
     "Ultra-photorealistic cinematic epic fantasy, wide 16:9 composition. A broken wooden suspension bridge spans a deep river gorge under an approaching storm. Wet boards, a frayed guide rope and one dangerous open gap are clearly visible. [PROTAGONIST] stands at the near edge with several other letter-bearing travelers, all studying the same crossing. Compose the scene so the four possible responses are visually available without showing four separate stories: a coil of rope near the protagonist, visible stone and timber supports beneath the bridge, loose boards and cord that could reinforce the gap, and wind moving through hanging cloth as a clue to timing. The far bank and the single continuing road are visible; there is no alternate trail. Emotional focus: courage, analysis, invention and patience as equally dignified ways of meeting fear. Dramatic gray sky, rain beginning, realistic height and scale, restrained danger rather than action spectacle. No one has fallen, no combat, no fantasy monsters, no split screen, no text, no logos."),
    (4, "La tormenta en la torre de vigilancia", "Convertir la espera en una prueba psicológica sin alterar el destino del grupo.",
     "Ultra-photorealistic cinematic epic fantasy, wide 16:9 interior. All travelers have crossed the bridge and now shelter inside a ruined stone watchtower during violent rain. [PROTAGONIST] is among them, neither isolated nor centered as a heroic savior. Water enters through one broken window; old route markings cover a wall; damp packs and supplies lie across the floor; a small lantern lights worried faces. The composition should contain several meaningful possibilities at once: dry space that could be organized, the broken window that could be covered with cloth and timber, ancient markings that could be studied, and a small group ready to listen to a calming story. Outside, lightning briefly reveals the road continuing toward the coast. Emotional focus: how a person inhabits uncertainty when the storm cannot be defeated, only endured. Cool slate rain and warm lantern contrast, intimate faces, realistic wet clothing and masonry. No branching exits, no argument turning violent, no supernatural event, no readable writing, no text or logos."),
    (5, "Una desconocida junto al fuego", "Crear intimidad y empatía sin convertir a Mara en una compañera permanente.",
     "Ultra-photorealistic cinematic epic fantasy, wide 16:9 composition. Quiet dusk at a forest shelter after the storm. [PROTAGONIST] and MARA sit on opposite sides of a modest campfire, close enough to speak honestly but each with a separate pack ready for a different pace at dawn. Two torn maps lie between them on a flat stone; both point to the same lighthouse, but the coastlines differ. Mara has auburn hair, a deep green shawl and a weathered traveler’s coat. Her expression is perceptive, calm and self-contained. Include small visual invitations to compare maps, redraw a route, share a personal reason or laugh together, but freeze the moment before any single choice is made. Blue evening forest surrounds a warm circle of firelight. Emotional focus: two unfinished inner journeys briefly recognizing one another without possession, rescue or obligation. Realistic faces, hands, paper and fire. No romance pose, no third companion, no promise to travel together, no magical blue flames, no readable text, no logos."),
    (6, "La carreta inmóvil", "Mostrar cooperación práctica y conflicto de método, no una ruta alternativa.",
     "Ultra-photorealistic cinematic epic fantasy, wide 16:9 composition. Midday on a narrow mountain pass. A sturdy supply cart has sunk axle-deep into thick mud. Its cargo is clearly practical for the coastal camp and lighthouse: sealed lamp-oil jars, folded blankets and sacks of food. [PROTAGONIST], the driver and several travelers gather around one trapped wheel. The frame should naturally show all usable resources: a long timber suitable as a lever, removable cargo, rope, open positions for coordinated pushing and several people offering competing ideas. The road ahead remains visible, descending toward the same coast and distant lighthouse; turning back is not framed as a separate story. Emotional focus: what someone contributes when everyone agrees on the goal but not on the method. Bright overcast light after rain, mud on boots and hems, realistic exertion and body language, grounded medieval transport. No overturned cart, no combat, no side road, no new permanent companion, no text or logos."),
    (7, "La criatura bajo el cobertizo", "Explorar cuidado, observación, diseño y regulación colectiva ante una criatura asustada.",
     "Ultra-photorealistic cinematic epic fantasy, wide 16:9 composition. Soft rain in an ancient cedar forest. Beneath a partially collapsed roadside shelter, a young griffin is trapped by a leather strap twisted around a fallen branch. Its soaked tawny-and-gray feathers cling naturally to its body, one wing trembles and its expression is frightened rather than aggressive. [PROTAGONIST] and a few travelers keep respectful distance. Place a bowl of water, cloth, loose branches and a clear potential escape corridor within the composition, while the exact action has not yet been chosen. The strap and the branch holding its tension must be visually understandable. Emotional focus: the moral weight of approaching a vulnerable being that cannot know your intention. Diffused green forest light, rain droplets, realistic animal anatomy combining eagle and lion, believable scale slightly larger than a wolf. No attack, no blood, no rider, no griffin army, no glowing magic, no text or logos."),
    (8, "La costa bajo la luna", "Presentar el faro en toda su escala y preparar de forma creíble la creciente.",
     "Ultra-photorealistic cinematic epic fantasy, panoramic wide 16:9 composition. At sunset, the cedar forest opens onto a vast rugged coast. For the first time, the Lighthouse of Vigil is fully visible on a black-rock island across a narrow channel, tall and ancient, its lantern currently dark. A full moon rises opposite the last orange light. In the middle ground, an elevated fishing camp occupies a safe stone terrace while low unloading platforms and stairs descend toward the water. [PROTAGONIST] arrives with other travelers as fishers secure boats, lift belongings and study the changing wind. Show high ground, exits, shared food and the vulnerable lower route in one coherent scene. Emotional focus: awe at reaching the destination’s edge, followed by the quiet recognition that nature sets its own terms. Cinematic sunset, silver-blue sea, realistic tents, nets and stonework. No flooding yet, no sea monster, no fantasy city, no alternate lighthouse, no text or logos."),
    (9, "La creciente", "Crear presión realista donde todos alcanzan el mismo lugar seguro.",
     "Ultra-photorealistic cinematic epic fantasy, wide 16:9 night composition. The full-moon tide has risen over the fishing camp’s lower unloading platform and covered the lowest stair. The elevated stone terrace remains clearly safe. Several travelers are still climbing with children, soaked packs and supply crates. A guide rope stretches across the exposed edge. [PROTAGONIST] stands where help is most needed while people form an emerging relay; the image should also reveal the narrow passage where human traffic and cargo could be separated. Across the silver channel, the Lighthouse of Vigil waits on its island. Emotional focus: urgency without panic, and the psychological choice of where to place oneself when no individual can stop the sea. Strong moonlight, warm camp lanterns, realistic waves and wet stone, credible human scale. Everyone is moving toward the same high ground. No casualties, no collapsing cliff, no combat, no supernatural wave, no text or logos."),
    (10, "Los cuatro objetos", "Representar la tentación como una proyección de necesidad, con los cuatro objetos en una sola imagen.",
     "Ultra-photorealistic cinematic epic fantasy, wide 16:9 composition. Inside a small chamber carved into a sea cliff, [PROTAGONIST] stands before four separate stone plinths arranged in a shallow arc. Each holds one intriguing, materially distinct object: a narrow dark-steel dagger, an old iron key with an unfamiliar tooth pattern, a clear glass vial containing amber liquid, and a weathered charcoal travel cloak with visible repairs. Moonlight enters from the open sea passage while a low hidden fire adds warm highlights. The protagonist’s hand is still lowered; no object has been chosen. An inscription exists as worn abstract marks in the stone but must not be readable. Emotional focus: the revealing pause in which need, control, curiosity and protection compete before action. Precise object detail, realistic metal, glass, wool and rock, restrained mystical atmosphere. Exactly four objects, no crown, no treasure pile, no weapon combat, no text, no logos."),
    (11, "El guardián del sendero", "Hacer visible que el objeto debe soltarse y que todas las opciones conducen al mismo ascenso.",
     "Ultra-photorealistic cinematic epic fantasy, wide 16:9 moonlit composition. Beyond the sea-cliff chamber, a dignified faun guardian waits beside a flat ritual stone at the base of the final lighthouse ascent. He has natural tawny fur, subtle curved horns, intelligent human features, goat legs, a worn walking staff and a warm brass lantern. He is solemn but not menacing. [PROTAGONIST] stands opposite him holding [CHOSEN OBJECT] with both the reluctance and understanding of someone being asked to release what felt necessary. The flat stone is empty and clearly ready to receive the object. Behind the faun, one narrow stair climbs toward the Lighthouse of Vigil; no alternate path exists. Other travelers remain softly out of focus nearby so a shared gesture is possible without changing the group. Emotional focus: letting go as a form of self-knowledge. Silver moonlight, warm lantern, realistic anatomy and materials. No payment scene, no combat stance, no magical portal, no multiple objects, no text or logos."),
    (12, "El banquete de quienes llegaron", "Ofrecer descanso, pertenencia y humanidad antes del trabajo final.",
     "Ultra-photorealistic cinematic epic fantasy, wide 16:9 interior. Inside the base of the ancient lighthouse, a long wooden table holds simple bread, steaming broth, fruit and water for travelers who have arrived from the same road. [PROTAGONIST] enters the warm room. Familiar faces from the bridge, watchtower and coastal camp sit among strangers; Mara may be visible at a distance, but she is not staged as the protagonist’s companion. Wet cloaks dry near a hearth. Some seats can be moved, one place is open, and several people are ready to exchange stories or laughter. The meal feels generous but modest, with no hierarchy at the head of the table. Emotional focus: the difficulty and relief of accepting belonging without having to prove that one deserves it. Warm candle and hearth light, realistic food, tired faces and old stone. No royal feast, no alcohol excess, no magical food, no conflict, no readable banners, no text or logos."),
    (13, "El faro apagado", "Mostrar que el faro vuelve a encenderse mediante esfuerzos distintos que convergen.",
     "Ultra-photorealistic cinematic epic fantasy, wide 16:9 interior with vertical depth. The maintenance chamber near the top of the Lighthouse of Vigil is dark except for work lanterns. [PROTAGONIST] and the other travelers face one shared task: restore the main lantern before ships cross the reef. Clearly show heavy sealed oil canisters beside the stair, a bent brass lamp support on a workbench, an unfinished duty schedule represented by blank marks and tokens rather than readable text, and a wall of small polished reflectors that can multiply one flame. Different people are already preparing to help, but the decisive work has not yet begun. Through arched windows, distant ship lights wait beyond black rocks. Emotional focus: purpose becoming collective, with strength, invention, coordination and transformation equally necessary. Brass, glass, oil, rope and stone rendered realistically; dramatic low light. No modern machinery, no electricity, no explosion, no single heroic pose, no text or logos."),
    (14, "Lo que dejas para el próximo viajero", "Crear un momento íntimo de legado antes de la ceremonia final.",
     "Ultra-photorealistic cinematic epic fantasy, wide 16:9 intimate composition. The lighthouse lantern now burns above, sending a broad warm beam through a window and across the moonlit sea. At a worn wooden desk, [PROTAGONIST] looks at an old travelers’ notebook opened to its final blank page. Earlier pages contain small hand-drawn maps, abstract symbols, sketches of the bridge and short marks that are not readable language. A charcoal pencil rests beside the book. The room carries quiet traces of the completed road: damp rope, a folded cloth, a single cedar leaf and reflected lighthouse light, but none should look like collectible trophies. The protagonist pauses before writing, aware that the next traveler is unknown. Emotional focus: choosing what to pass forward without choosing another person’s path for them. Warm beacon light, cool full moon, tactile paper and wood, contemplative close framing. No visible future traveler, no readable words, no magical writing, no text overlay, no logos."),
    (15, "La melodía que parecía recordarte", "Cerrar el ciclo con el flautista, la hipnosis y la sensación de haber vivido este sueño antes.",
     "Ultra-photorealistic cinematic epic fantasy, wide 16:9 composition. At midnight inside a circular chamber open to the full moon, the masked flutist stands at the center of a ring of travelers and raises a slender silver flute. His mask is simple ivory with one narrow obsidian line, his layered coat is deep indigo and his presence is calm, unreadable and magnetic. [PROTAGONIST] stands among the travelers with the open doorway still visible behind, allowing a clear sense of personal space and choice. Some people watch the flute, others watch one another; no one is unconscious yet. Above, the restored lighthouse burns, and its rotating beam crosses the chamber as the first faint visual suggestion of music appears only through drifting dust and circular light, never literal notes. Emotional focus: recognition without explanation, danger without violence, and the sensation that the melody already knows the listener. Realistic moonlit faces, cloth and stone, restrained hypnotic atmosphere. No concert stage, no glowing eyes, no mind-control beams, no readable text, no logos."),
]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement('w:tblHeader')
    tbl_header.set(qn('w:val'), 'true')
    tr_pr.append(tbl_header)


def suppress_paragraph_borders(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    existing = p_pr.find(qn('w:pBdr'))
    if existing is not None:
        p_pr.remove(existing)
    borders = OxmlElement('w:pBdr')
    for edge in ('top', 'left', 'bottom', 'right', 'between', 'bar'):
        border = OxmlElement(f'w:{edge}')
        border.set(qn('w:val'), 'nil')
        borders.append(border)
    p_pr.append(borders)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.72)
section.left_margin = Inches(0.82)
section.right_margin = Inches(0.82)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal'].font.size = Pt(10.5)
styles['Normal'].font.color.rgb = RGBColor(35, 39, 37)
styles['Normal'].paragraph_format.space_after = Pt(7)
styles['Normal'].paragraph_format.line_spacing = 1.13

for style_name, size in [('Title', 28), ('Heading 1', 20), ('Heading 2', 15), ('Heading 3', 11)]:
    style = styles[style_name]
    style.font.name = 'Georgia'
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.bold = style_name == 'Heading 3'
    style.paragraph_format.keep_with_next = True
    style.paragraph_format.space_before = Pt(15 if style_name != 'Title' else 0)
    style.paragraph_format.space_after = Pt(7)

if 'Prompt' not in [s.name for s in styles]:
    prompt_style = styles.add_style('Prompt', WD_STYLE_TYPE.PARAGRAPH)
else:
    prompt_style = styles['Prompt']
prompt_style.font.name = 'Aptos'
prompt_style.font.size = Pt(9.7)
prompt_style.font.color.rgb = RGBColor(28, 35, 31)
prompt_style.paragraph_format.left_indent = Inches(0.18)
prompt_style.paragraph_format.right_indent = Inches(0.12)
prompt_style.paragraph_format.space_after = Pt(10)
prompt_style.paragraph_format.line_spacing = 1.08

title = doc.add_paragraph(style='Title')
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run('Prompts visuales para The Unwritten Road')
suppress_paragraph_borders(title)
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run('Guía de continuidad para GPT Image 2')
run.bold = True
run.font.size = Pt(13)
run.font.color.rgb = RGBColor(48, 67, 57)
meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.add_run('MVP 1.1  ·  15 escenas  ·  Formato panorámico 16:9').italic = True

doc.add_paragraph()
intro = doc.add_paragraph()
intro.add_run('Propósito. ').bold = True
intro.add_run('Este documento permite generar una imagen por momento narrativo sin crear rutas visuales paralelas. Todas las decisiones ocurren dentro de la misma composición y conducen al siguiente punto de la historia. La variación psicológica pertenece a la respuesta del usuario, no al orden de las escenas.')

usage = doc.add_paragraph()
usage.add_run('Cómo usarlo. ').bold = True
usage.add_run('En cada generación, copia primero el bloque maestro, reemplaza las variables indicadas y añade el prompt de la escena. Conserva la misma descripción del protagonista, Mara, el fauno, el flautista y el faro en todas las imágenes. Usa la primera imagen aprobada del protagonista como referencia visual en las siguientes generaciones cuando la interfaz lo permita.')

doc.add_heading('Bloque maestro de continuidad', level=1)
master = (
    "Create one cinematic story frame from the same grounded epic-fantasy world as the previous images. Ultra-photorealistic environmental storytelling, emotionally intimate rather than action-driven, believable medieval materials, natural human anatomy, subtle atmospheric depth, fine film grain, physically plausible light, premium fantasy landscape photography, wide 16:9 composition with room for website interface overlays. Maintain exact continuity of faces, clothing, props, architecture, weather progression and travel wear. The Lighthouse of Vigil is a tall weathered pale-stone lighthouse built on a black-rock island, with an old bronze lantern room. The protagonist always wears an earth-brown tunic, deep moss-green travel cloak, dark leather boots and the same weathered leather backpack. Avoid generic video-game concept art, glossy armor, excessive magic, exaggerated heroic poses, split screens, collages, duplicated characters, modern objects, readable writing, captions, UI, logos and watermarks. Show only the single canonical moment described below."
)
p = doc.add_paragraph(master, style='Prompt')

doc.add_heading('Variables de personaje', level=1)
table = doc.add_table(rows=1, cols=2)
table.style = 'Table Grid'
table.autofit = False
table.columns[0].width = Inches(1.55)
table.columns[1].width = Inches(5.25)
hdr = table.rows[0]
set_repeat_table_header(hdr)
for cell, text in zip(hdr.cells, ['Variable', 'Texto para reemplazar']):
    set_cell_shading(cell, '263D33')
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    run = cell.paragraphs[0].add_run(text)
    run.bold = True
    run.font.color.rgb = RGBColor(255, 255, 255)

variables = [
    ('[PROTAGONIST] hombre', 'The same male traveler, around 30 years old, medium olive skin, wavy dark-brown hair to the neck, thoughtful brown eyes, lean build, no visible armor.'),
    ('[PROTAGONIST] mujer', 'The same female traveler, around 30 years old, medium olive skin, long dark-brown hair in one practical braid, thoughtful brown eyes, lean build, no visible armor.'),
    ('MARA', 'The same independent female traveler in her early thirties, light olive skin, auburn hair loosely tied back, deep green shawl, weathered gray-brown coat, calm perceptive expression.'),
    ('[CHOSEN OBJECT]', 'Use only the object selected in scene 10: the narrow dagger, the old iron key, the amber vial or the weathered charcoal cloak.'),
]
for i, (key, value) in enumerate(variables):
    cells = table.add_row().cells
    if i % 2:
        for cell in cells:
            set_cell_shading(cell, 'F3F5F3')
    cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cells[0].paragraphs[0].add_run(key).bold = True
    cells[1].paragraphs[0].add_run(value)

doc.add_heading('Reglas para mantener un solo hilo', level=1)
for text in [
    'Genera una sola composición por escena. No uses cuadrículas, antes y después ni cuatro versiones dentro de una imagen.',
    'Muestra recursos que permitan imaginar las cuatro respuestas, pero congela la acción antes de que una de ellas altere la escena.',
    'No añadas compañeros permanentes, objetos coleccionables, rutas alternas, combates ni consecuencias que no aparezcan en el siguiente prompt.',
    'Conserva el desgaste acumulado del protagonista y la progresión temporal desde el amanecer hasta la luna llena.',
    'Evita texto legible dentro de cartas, mapas, inscripciones y cuadernos para reducir errores visuales.',
]:
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(text)

doc.add_heading('Secuencia visual y prompts por escena', level=1)
current_act = None
act_names = {
    1: 'Acto 1 El llamado',
    2: 'Acto 2 El umbral',
    3: 'Acto 3 Aliados y desconocidos',
    4: 'Acto 4 Las pruebas',
    5: 'Acto 5 La ofrenda',
    6: 'Acto 6 El faro',
    7: 'Acto 7 El retorno',
}
act_by_scene = {1:1, 2:1, 3:2, 4:2, 5:3, 6:3, 7:4, 8:4, 9:4, 10:5, 11:5, 12:6, 13:6, 14:7, 15:7}

for number, name, purpose, prompt in SCENES:
    act = act_by_scene[number]
    if act != current_act:
        heading = doc.add_heading(act_names[act], level=1)
        if current_act is not None:
            heading.paragraph_format.page_break_before = True
        current_act = act
    doc.add_heading(f'Escena {number:02d} {name}', level=2)
    p = doc.add_paragraph()
    p.add_run('Objetivo narrativo  ').bold = True
    p.add_run(purpose)
    label = doc.add_paragraph()
    label.paragraph_format.keep_with_next = True
    run = label.add_run('PROMPT PARA COPIAR')
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(75, 91, 82)
    doc.add_paragraph(prompt, style='Prompt')

doc.add_heading('Control final antes de aprobar una imagen', level=1)
for text in [
    'El protagonista conserva rostro, cabello, capa, mochila y proporciones.',
    'La imagen representa el momento previo a la decisión y no una consecuencia exclusiva.',
    'El faro, el clima y la hora coinciden con la escena anterior y la siguiente.',
    'No aparecen texto legible, interfaces, logotipos, marcas de agua ni elementos modernos.',
    'La emoción principal se entiende aun sin leer la historia.',
]:
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(text)

for section in doc.sections:
    header = section.header.paragraphs[0]
    header.text = 'THE UNWRITTEN ROAD  ·  GUÍA VISUAL'
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header.runs[0].font.name = 'Aptos'
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor(95, 105, 99)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run('Documento de producción visual  ·  MVP 1.1')
    run.font.name = 'Aptos'
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(95, 105, 99)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
